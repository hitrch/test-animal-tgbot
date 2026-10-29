import io
import shutil

import docx
import pytest

from bot.archive import ARCHIVE_DIR
from bot.extract import UnsupportedFormatError, extract_text


def test_txt_in_utf8_with_windows_notepad_bom():
    assert extract_text("Порада.txt", "﻿Внесок — 4 жолуді\n".encode()) == "Внесок — 4 жолуді"


def test_txt_in_old_windows_encoding():
    assert extract_text("Порада.txt", "Внесок — 4 жолуді".encode("cp1251")) == "Внесок — 4 жолуді"


def test_docx_paragraphs_and_tables():
    document = docx.Document()
    document.add_paragraph("Членські внески")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Їжак"
    table.rows[0].cells[1].text = "4 жолуді"
    data = io.BytesIO()
    document.save(data)

    assert extract_text("Внески.docx", data.getvalue()) == "Членські внески\nЇжак | 4 жолуді"


def test_pdf_with_a_text_layer():
    path = ARCHIVE_DIR / "Броди і коли бобри відкривають греблю.pdf"

    text = extract_text(path.name, path.read_bytes())

    assert "Бобри відкривають греблю двічі на тиждень" in text


# Tesseract is installed in the Docker image; on a developer machine it may be missing
@pytest.mark.skipif(shutil.which("tesseract") is None, reason="Tesseract is not installed")
def test_scanned_pdf_is_recognized_with_ocr():
    path = ARCHIVE_DIR / "Грибний путівник хвойного лісу.pdf"

    text = extract_text(path.name, path.read_bytes())

    assert "гриб" in text.lower()


def test_other_formats_are_not_supported():
    with pytest.raises(UnsupportedFormatError):
        extract_text("Таблиця.xlsx", b"...")