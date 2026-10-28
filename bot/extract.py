"""Turn an uploaded or archive file into plain text."""
import io
from pathlib import Path

import docx
import pypdf
import pytesseract
from PIL import Image, ImageOps

SUPPORTED_EXTENSIONS = (".txt", ".docx", ".pdf")

# Pixels lighter than this become white. Chosen by experiments on the archive scans:
# it removes the birch-bark background but keeps the dark text.
OCR_THRESHOLD = 90
# "eng" helps to tell the digit 3 from the Ukrainian letter З
OCR_LANGUAGES = "ukr+eng"


class UnsupportedFormatError(Exception):
    pass


def extract_text(filename: str, data: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension == ".txt":
        return _read_txt(data)
    if extension == ".docx":
        return _read_docx(data)
    if extension == ".pdf":
        return _read_pdf(data)
    raise UnsupportedFormatError(extension)


def _read_txt(data: bytes) -> str:
    # utf-8-sig also strips the BOM that Windows Notepad adds
    try:
        return data.decode("utf-8-sig").strip()
    except UnicodeDecodeError:
        return data.decode("cp1251").strip()


def _read_docx(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines).strip()


def _read_pdf(data: bytes) -> str:
    reader = pypdf.PdfReader(io.BytesIO(data))
    text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if text:
        return text
    # No text layer means the PDF is a scan: recognize the page images
    return _ocr_pdf(reader)


def _ocr_pdf(reader: pypdf.PdfReader) -> str:
    parts = []
    for page in reader.pages:
        for image in page.images:
            parts.append(pytesseract.image_to_string(_to_black_and_white(image.image), lang=OCR_LANGUAGES))
    return "\n".join(parts).strip()


def _to_black_and_white(image: Image.Image) -> Image.Image:
    gray = ImageOps.grayscale(image)
    return gray.point(lambda pixel: 255 if pixel > OCR_THRESHOLD else 0)